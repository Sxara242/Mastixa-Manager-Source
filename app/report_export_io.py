"""Small file/worksheet helpers for the annual and data-quality exports."""
from contextlib import contextmanager
import os
from pathlib import Path
import tempfile

from .language import tr


@contextmanager
def export_destination(path: Path):
    """Promote a completed sibling file; preserve an existing file on failure."""
    fd, name = tempfile.mkstemp(prefix=f".{path.stem}-", suffix=path.suffix, dir=path.parent)
    os.close(fd)
    staging = Path(name)
    try:
        yield staging
        if not staging.stat().st_size:
            raise OSError(tr("Η εξαγωγή παρήγαγε κενό αρχείο."))
        os.replace(staging, path)
    finally:
        staging.unlink(missing_ok=True)


def write_sheet(workbook, title, headers, rows, widths):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    sheet = workbook.create_sheet(title)
    for values in (headers, *rows):
        row_number = sheet.max_row + 1 if sheet.cell(1, 1).value is not None else 1
        for column, value in enumerate(values, 1):
            if isinstance(value, str) and len(value) > 32767:
                raise ValueError(tr("Το κείμενο υπερβαίνει το όριο κελιού Excel (32767 χαρακτήρες)."))
            cell = sheet.cell(row_number, column, value)
            if isinstance(value, str):
                # Raw names/issues beginning with '=' are text, never formulas.
                cell.data_type = "s"
            elif isinstance(value, float):
                cell.number_format = "#,##0.###"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="315B4C")
    sheet.row_dimensions[1].height = 32
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    return sheet


def export_issue_xlsx(path: Path, headers, rows, title: str) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    try:
        write_sheet(workbook, title, headers, rows, (18, 26, 28, 65, 65))
        with export_destination(path) as staging:
            workbook.save(staging)
    finally:
        workbook.close()
