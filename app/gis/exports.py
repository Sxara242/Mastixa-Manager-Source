"""Coordinate exports use one ordered, validated snapshot; no database writes."""
from __future__ import annotations
import csv
import io
import json
import re
from dataclasses import dataclass
from decimal import Decimal
from html import escape
from pathlib import Path
from xml.etree import ElementTree as ET
from .geometry import normalize,vertices,parts


@dataclass(frozen=True)
class ExportParcel:
    identity: str
    name: str
    kaek: str
    source: str
    geometry: object


def snapshot(records):
    result=[];seen=set();count=0
    for row in records:
        if row['id'] in seen:continue
        seen.add(row['id']);geometry=normalize(json.loads(row['original_geojson']),row['source_crs'])
        count+=geometry.vertex_count
        if count>100000:raise ValueError('Export exceeds 100,000 vertices; select fewer fields')
        result.append(ExportParcel(row['id'],row.get('current_name',row['name']),row.get('current_kaek',row['kaek']),row['geometry_source'],geometry))
    if not result:raise ValueError('Select at least one parcel with saved geometry')
    return result


def filename(parcels,mode,extension):
    title=parcels[0].name+'_KAEK_'+parcels[0].kaek if len(parcels)==1 else 'Mastixa_'+str(len(parcels))+'_fields'
    title=re.sub(r'[<>:"/\\|?*\x00-\x1f]','_',title).strip(' .')[:100] or 'Mastixa'
    if title.split('.')[0].upper() in {'CON','PRN','AUX','NUL',*[f'COM{i}' for i in range(1,10)],*[f'LPT{i}' for i in range(1,10)]}:title='_'+title
    if extension not in ('csv','xlsx','pdf','geojson','kml'):raise ValueError('Unsupported export format')
    return f'{title}_coordinates_{"WGS84" if mode=="wgs84" else "source"}.{extension}'


def table(parcels,mode):
    if mode not in ('wgs84','source'):raise ValueError('Unsupported coordinate mode')
    axes=['Latitude','Longitude'] if mode=='wgs84' else ['X / Longitude','Y / Latitude']
    rows=[['Field Name','KAEK','Coordinate System','EPSG','Area (m²)','Perimeter (m)','Vertex count','Part','Ring','Vertex',*axes]]
    for parcel in parcels:
        g=parcel.geometry;crs='EPSG:4326' if mode=='wgs84' else g.source_crs
        coordinates=g.wgs84 if mode=='wgs84' else g.original
        for part,ring,index,x,y in vertices(coordinates):
            a,b=(y,x) if mode=='wgs84' else (x,y)
            rows.append([parcel.name,parcel.kaek,'WGS84 / GPS' if mode=='wgs84' else 'Original source XY',crs,
                         Decimal(f'{g.area_m2:.2f}'),Decimal(f'{g.perimeter_m:.2f}'),g.vertex_count,part,ring,index,
                         Decimal(f'{a:.8f}'),Decimal(f'{b:.8f}')])
    return rows


def csv_bytes(rows):
    out=io.StringIO(newline='');writer=csv.writer(out,delimiter=';',lineterminator='\r\n')
    for row in rows:
        writer.writerow(["'"+cell if isinstance(cell,str) and cell.lstrip()[:1] in ('=','+','-','@') else cell for cell in row])
    return out.getvalue().encode('utf-8-sig')


def xlsx_bytes(rows):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font,PatternFill
    from openpyxl.utils import get_column_letter
    workbook=Workbook(write_only=True);sheet=workbook.create_sheet('Όλες οι κορυφές');sheet.freeze_panes='A2'
    for i,width in enumerate([30,20,25,18,18,18,15,9,9,10,23,23],1):sheet.column_dimensions[get_column_letter(i)].width=width
    for index,row in enumerate(rows):
        cells=[]
        for value in row:
            cell=WriteOnlyCell(sheet,value=value)
            if isinstance(value,str):cell.data_type='s'
            if isinstance(value,Decimal):cell.number_format='0.00000000' if index and len(cells)>=10 else '0.00'
            if index==0:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='23683E')
            cells.append(cell)
        sheet.append(cells)
    out=io.BytesIO();workbook.save(out);return out.getvalue()


def properties(parcel):
    g=parcel.geometry
    return dict(field_id=parcel.identity,field_name=parcel.name,kaek=parcel.kaek,geometry_source=parcel.source,
                source_crs=g.source_crs,coordinate_system='WGS84',epsg='EPSG:4326',area_m2=g.area_m2,
                perimeter_m=g.perimeter_m,vertex_count=g.vertex_count)


def geojson_bytes(parcels):
    return json.dumps(dict(type='FeatureCollection',features=[dict(type='Feature',id=p.identity,properties=properties(p),geometry=p.geometry.wgs84) for p in parcels]),ensure_ascii=False,allow_nan=False).encode('utf-8')


def kml_bytes(parcels):
    namespace='http://www.opengis.net/kml/2.2';ET.register_namespace('',namespace)
    def tag(name):return '{'+namespace+'}'+name
    root=ET.Element(tag('kml'));document=ET.SubElement(root,tag('Document'))
    for parcel in parcels:
        placemark=ET.SubElement(document,tag('Placemark'));ET.SubElement(placemark,tag('name')).text=parcel.name
        extended=ET.SubElement(placemark,tag('ExtendedData'))
        for key,value in properties(parcel).items():ET.SubElement(ET.SubElement(extended,tag('Data'),name=key),tag('value')).text=str(value)
        polygons=list(parts(parcel.geometry.wgs84));parent=ET.SubElement(placemark,tag('MultiGeometry')) if len(polygons)>1 else placemark
        for rings in polygons:
            polygon=ET.SubElement(parent,tag('Polygon'))
            for index,ring in enumerate(rings):
                boundary=ET.SubElement(polygon,tag('outerBoundaryIs' if index==0 else 'innerBoundaryIs'))
                linear=ET.SubElement(boundary,tag('LinearRing'));ET.SubElement(linear,tag('coordinates')).text=' '.join(f'{lon:.8f},{lat:.8f},0' for lon,lat in ring)
    return ET.tostring(root,encoding='utf-8',xml_declaration=True)


def pdf_file(path,parcels,mode):
    from PySide6.QtCore import QMarginsF
    from PySide6.QtGui import QFont,QPageLayout,QPageSize,QTextDocument
    from PySide6.QtPrintSupport import QPrinter
    printer=QPrinter(QPrinter.PrinterMode.HighResolution);printer.setOutputFormat(QPrinter.OutputFormat.PdfFormat);printer.setOutputFileName(str(path))
    printer.setPageSize(QPageSize(QPageSize.PageSizeId.A4));printer.setPageMargins(QMarginsF(14,14,14,14),QPageLayout.Unit.Millimeter)
    sections=[]
    for parcel in parcels:
        rows=table([parcel],mode);g=parcel.geometry
        heading=f'<h2>{escape(parcel.name)}</h2><p>KAEK: {escape(parcel.kaek)}<br>{escape(rows[1][2])} · {escape(rows[1][3])}<br>Area: {g.area_m2:.2f} m² · Perimeter: {g.perimeter_m:.2f} m · Vertices: {g.vertex_count}<br>Source: {escape(parcel.source)} · Source CRS: {escape(g.source_crs)}</p>'
        headers=''.join(f'<th>{escape(str(value))}</th>' for value in rows[0][7:])
        body=''.join('<tr>'+''.join(f'<td>{escape(str(value))}</td>' for value in row[7:])+'</tr>' for row in rows[1:])
        sections.append(heading+f'<table width="100%" cellspacing="0" cellpadding="5"><thead><tr>{headers}</tr></thead>{body}</table>')
    doc=QTextDocument();doc.setDefaultFont(QFont('Segoe UI',10))
    doc.setHtml('<html><head><style>body{color:#20372b}th{background:#e2eee5}td{border-bottom:1px solid #ddd}h2{font-size:15pt}</style></head><body>'+ '<div style="page-break-before:always"></div>'.join(sections)+'</body></html>');doc.print_(printer)
    if not Path(path).exists() or Path(path).stat().st_size==0:raise OSError('PDF output failed')


def write(path,parcels,mode,extension):
    if extension=='pdf':pdf_file(path,parcels,mode);return
    if extension=='csv':data=csv_bytes(table(parcels,mode))
    elif extension=='xlsx':data=xlsx_bytes(table(parcels,mode))
    elif extension=='geojson':data=geojson_bytes(parcels)
    elif extension=='kml':data=kml_bytes(parcels)
    else:raise ValueError('Unsupported export format')
    Path(path).write_bytes(data)
