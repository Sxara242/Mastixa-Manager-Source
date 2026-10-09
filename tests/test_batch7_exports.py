"""Windows export contracts: shared raw snapshots, safe files and EL/EN output."""
import csv
from dataclasses import FrozenInstanceError
from pathlib import Path
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from PySide6.QtWidgets import QApplication
from PySide6.QtGui import QFontDatabase

from app.annual_report import AnnualFarmReportPage
from app import annual_report_exports as exports
from app.data_quality import DataQualityPage, _IssueText
from tests import test_report_quantity_contract as contract


class Batch7ExportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.qt = QApplication.instance() or QApplication([])
        # Windows' offscreen plugin has no system font database. Match the
        # existing coordinate-PDF test's explicit Greek-capable font setup.
        font = Path("C:/Windows/Fonts/segoeui.ttf")
        if font.is_file():
            QFontDatabase.addApplicationFont(str(font))

    # Reuse the existing mixed-unit/two-year fixture, without collecting its tests.
    setUp = contract.ReportQuantityContractTests.setUp
    page = contract.ReportQuantityContractTests.page
    select = contract.ReportQuantityContractTests.select
    snapshot = contract.ReportQuantityContractTests.snapshot

    def annual(self, product=None, year=2027):
        page = self.page(AnnualFarmReportPage)
        self.select(page.year_filter, year)
        self.select(page.product_filter, product)
        return page

    def workbook(self, path):
        workbook = load_workbook(path)
        self.addCleanup(workbook.close)
        return workbook

    def test_control_pdf_default_and_live_el_en_el(self):
        page = self.annual(self.pieces)
        before = self.snapshot()
        self.assertEqual(["pdf", "xlsx", "csv"], list(page.export_actions))
        self.assertIs(page.export_button.defaultAction(), page.export_actions["pdf"])
        for code in ("el", "en", "el"):
            self.language.set_language(code, persist=False)
            self.language.apply_to(page)
            for kind, action in page.export_actions.items():
                self.assertEqual((kind.upper() + " export") if code == "en" else ("Εξαγωγή " + kind.upper()), action.text())
            self.assertEqual(2027, page.year_filter.currentData())
            self.assertEqual(self.pieces, page.product_filter.currentData())
        self.assertEqual(before, self.snapshot())

    def test_snapshot_immutable_independent_of_rendered_labels(self):
        page = self.annual(self.kg)
        snapshot = page._annual_report_snapshot()
        before = self.snapshot()
        page.production_metric[1].setText("WRONG 999 pieces")
        page.scope_note.setText("WRONG")
        page.product_table.item(0, 0).setText("WRONG")
        self.assertIs(snapshot, page._annual_report_snapshot())
        self.assertEqual((2027, self.kg, 55, 5, 140, 60),
                         (snapshot.year, snapshot.product_id, snapshot.products[0].produced,
                          snapshot.products[0].sold, snapshot.products[0].stock, snapshot.sales_revenue))
        self.assertIsNone(snapshot.net_result)
        with self.assertRaises(FrozenInstanceError):
            snapshot.year = 2030
        self.assertEqual(before, self.snapshot())

    def test_every_ui_export_receives_identical_selected_snapshot(self):
        page = self.annual(self.pieces, 2026)
        snapshot = page._annual_report_snapshot()
        for kind in ("pdf", "xlsx", "csv"):
            path = self.root / ("selected." + kind)
            with patch("app.annual_report.QFileDialog.getSaveFileName", return_value=(str(path), "")), \
                 patch("app.annual_report.export_annual_" + kind) as writer, patch("app.annual_report._message"):
                getattr(page, "export_" + kind)()
                writer.assert_called_once_with(path, snapshot)
        self.assertEqual((10, 2, 8, "pieces"),
                         (snapshot.products[0].produced, snapshot.products[0].sold,
                          snapshot.products[0].stock, snapshot.products[0].unit))

    def test_pdf_real_output_and_semantics_in_both_languages(self):
        from PySide6.QtPdf import QPdfDocument
        page = self.annual(None, 2026)
        before = self.snapshot()
        for code in ("el", "en"):
            self.language.set_language(code, persist=False)
            snapshot = page._annual_report_snapshot()
            html = exports.annual_html(snapshot)
            self.assertIn("110 kg", html)
            self.assertIn("10 pieces", html)
            self.assertNotIn("120 kg", html)
            self.assertIn("90 kg", html)
            self.assertIn("8 pieces", html)
            self.assertIn("176 EUR", html)
            path = self.root / (code + ".pdf")
            exports.export_annual_pdf(path, snapshot)
            document = QPdfDocument()
            self.assertEqual(QPdfDocument.Error.None_, document.load(str(path)))
            text = " ".join(document.getAllText(i).text() for i in range(document.pageCount()))
            document.close()
            self.assertIn("2026", text)
            self.assertIn("Annual Farm Report" if code == "en" else "Ετήσια Αναφορά Εκμετάλλευσης", text)
            self.assertIn("Mastixa", text)

        self.select(page.product_filter, self.pieces)
        selected = exports.annual_html(page._annual_report_snapshot())
        self.assertIn("10 pieces", selected)
        self.assertNotIn("110 kg", selected)
        self.assertIsNone(page._annual_report_snapshot().net_result)
        self.assertEqual(before, self.snapshot())

    def test_xlsx_numeric_units_carryover_and_unallocated_finance(self):
        page = self.annual(self.kg)
        path = self.root / "annual.xlsx"
        exports.export_annual_xlsx(path, page._annual_report_snapshot())
        workbook = self.workbook(path)
        self.assertEqual(["Σύνοψη", "Προϊόντα", "Οικονομικά"], workbook.sheetnames)
        self.assertEqual(("Mastixa", 55, 5, 60, 12, 140, "kg"), tuple(c.value for c in workbook["Προϊόντα"][2]))
        self.assertEqual(["s", "n", "n", "n", "n", "n", "s"], [c.data_type for c in workbook["Προϊόντα"][2]])
        finance = list(workbook["Οικονομικά"].values)
        self.assertEqual(4, finance[3][1])
        self.assertIn("Μη κατανεμημένα", finance[3][2])
        result = next(row for row in workbook["Σύνοψη"].values if row[0] == "Καθαρό αποτέλεσμα")
        self.assertIsNone(result[2])
        for sheet in workbook:
            self.assertEqual("A2", sheet.freeze_panes)
            self.assertTrue(sheet.auto_filter.ref)
            self.assertFalse(any(c.data_type == "f" for row in sheet for c in row))

    def test_all_products_xlsx_and_raw_formula_like_names(self):
        raw = '=Αποθήκευση <b>{year}</b>; "Ω" ' + "Long product " * 12
        self.db.execute("UPDATE products SET name=?, unit=? WHERE id=?", (raw, "pieces {unit}", self.pieces))
        page = self.annual(None, 2026)
        before = self.snapshot()
        self.language.set_language("en", persist=False)
        path = self.root / "all.xlsx"
        exports.export_annual_xlsx(path, page._annual_report_snapshot())
        workbook = self.workbook(path)
        rows = list(workbook["Products"].values)
        row = next(row for row in rows[1:] if row[0] == raw)
        self.assertEqual((raw, 10, 2, 6, 3, 8, "pieces {unit}"), row)
        cell = next(row[0] for row in workbook["Products"] if row[0].value == raw)
        self.assertEqual("s", cell.data_type)
        result = next(row for row in workbook["Summary"].values if row[0] == "Net result")
        self.assertEqual(176, result[2])
        self.assertEqual(before, self.snapshot())

    def test_csv_preserves_existing_layout_bom_delimiter_and_values(self):
        page = self.annual(None, 2026)
        for code in ("el", "en"):
            self.language.set_language(code, persist=False)
            path = self.root / "annual.csv"
            exports.export_annual_csv(path, page._annual_report_snapshot())
            self.assertTrue(path.read_bytes().startswith(b"\xef\xbb\xbf"))
            with path.open(encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.reader(handle, delimiter=";"))
            self.assertEqual(7, len(rows))
            self.assertEqual("2026", rows[1][1])
            self.assertEqual("All products" if code == "en" else "Όλα τα προϊόντα", rows[2][1])
            self.assertEqual(["Mastixa", "110", "20", "200", "10", "90", "kg"], rows[5])
            self.assertEqual(["QA PRODUCT", "10", "2", "6", "3", "8", "pieces"], rows[6])

    def test_annual_cancel_creates_nothing(self):
        page = self.annual()
        before = set(self.root.iterdir())
        for kind in ("pdf", "xlsx", "csv"):
            with patch("app.annual_report.QFileDialog.getSaveFileName", return_value=("", "")), \
                 patch("app.annual_report.export_annual_" + kind) as writer, patch("app.annual_report._message") as message:
                getattr(page, "export_" + kind)()
                writer.assert_not_called()
                message.assert_not_called()
        self.assertEqual(before, set(self.root.iterdir()))

    def test_annual_extensions_and_actual_success_path(self):
        page = self.annual(self.pieces)
        for kind in ("pdf", "xlsx", "csv"):
            for suffix in ("", "." + kind.upper(), ".wrong"):
                path = self.root / (kind + suffix)
                expected = path if suffix == "." + kind.upper() else path.with_suffix("." + kind)
                with patch("app.annual_report.QFileDialog.getSaveFileName", return_value=(str(path), "")), \
                     patch("app.annual_report._message") as message:
                    getattr(page, "export_" + kind)()
                    self.assertTrue(expected.is_file())
                    self.assertEqual("information", message.call_args.args[1])
                    self.assertEqual(expected, message.call_args.kwargs["path"])

    def test_annual_failed_promotion_preserves_destination_and_cleans_staging(self):
        page = self.annual()
        for kind in ("pdf", "xlsx", "csv"):
            path = self.root / ("valid." + kind)
            path.write_bytes(b"previous valid file")
            before = set(self.root.iterdir())
            with patch("app.annual_report.QFileDialog.getSaveFileName", return_value=(str(path), "")), \
                 patch("app.report_export_io.os.replace", side_effect=OSError("Raw Αποθήκευση {path}")), \
                 patch("app.annual_report._message") as message:
                getattr(page, "export_" + kind)()
                self.assertEqual("critical", message.call_args.args[1])
                self.assertEqual("Raw Αποθήκευση {path}", message.call_args.kwargs["error"])
            self.assertEqual(b"previous valid file", path.read_bytes())
            self.assertEqual(before, set(self.root.iterdir()))

    def test_pdf_silent_print_failure_is_not_success(self):
        page = self.annual()
        path = self.root / "valid.pdf"
        path.write_bytes(b"previous PDF")
        before = set(self.root.iterdir())
        with patch("app.annual_report.QFileDialog.getSaveFileName", return_value=(str(path), "")), \
             patch.object(exports.QTextDocument, "print_"), patch("app.annual_report._message") as message:
            page.export_pdf()
        self.assertEqual("critical", message.call_args.args[1])
        self.assertEqual(b"previous PDF", path.read_bytes())
        self.assertEqual(before, set(self.root.iterdir()))

    def quality(self):
        from app.partner_links import ensure_partner_link_schema
        ensure_partner_link_schema(self.db)
        page = self.page(DataQualityPage)
        raw = '=Ω <raw> {path}; "Αποθήκευση"'
        page.all_issues = [dict(severity="WARNING", category="income", record=raw,
                              problem=_IssueText("Μη έγκυρη ημερομηνία: {value}.", value=raw),
                              fix="Συμπλήρωσε το ΚΑΕΚ αν είναι διαθέσιμο."),
                           dict(severity="ERROR", category="fields", record="excluded", problem="raw", fix="raw")]
        self.select(page.category_filter, "income")
        page._apply_filters()
        return page, raw

    def test_pdf_truncated_output_preserves_destination(self):
        path = self.root / "existing.pdf"
        path.write_bytes(b"previous valid PDF")
        before = set(self.root.iterdir())
        def truncated(printer):
            Path(printer.outputFileName()).write_bytes(b"%PDF-1.4\ntruncated document")
        with patch.object(exports.QTextDocument, "print_", side_effect=truncated), self.assertRaises(OSError):
            exports.export_annual_pdf(path, self.annual()._annual_report_snapshot())
        self.assertEqual(b"previous valid PDF", path.read_bytes())
        self.assertEqual(before, set(self.root.iterdir()))

    def test_owned_writer_errors_localized_and_preserve_existing_file(self):
        from app.report_export_io import export_destination, export_issue_xlsx
        page = self.annual()
        path = self.root / "existing.xlsx"
        path.write_bytes(b"existing valid destination")
        before = set(self.root.iterdir())
        for code in ("el", "en"):
            self.language.set_language(code, persist=False)
            with self.assertRaises(OSError) as caught:
                with export_destination(path):
                    pass
            self.assertEqual("The export produced an empty file." if code == "en" else
                             "Η εξαγωγή παρήγαγε κενό αρχείο.", str(caught.exception))
            with self.assertRaises(ValueError) as caught:
                export_issue_xlsx(path, ["raw header"], [("Ω" * 32768,)], "Issues")
            self.assertEqual("Text exceeds the Excel cell limit (32767 characters)." if code == "en" else
                             "Το κείμενο υπερβαίνει το όριο κελιού Excel (32767 χαρακτήρες).", str(caught.exception))
            with patch.object(exports.QTextDocument, "print_"), self.assertRaises(OSError) as caught:
                exports.export_annual_pdf(path, page._annual_report_snapshot())
            self.assertEqual("The PDF could not be written." if code == "en" else
                             "Δεν ήταν δυνατή η εγγραφή του PDF.", str(caught.exception))
            self.assertEqual(b"existing valid destination", path.read_bytes())
            self.assertEqual(before, set(self.root.iterdir()))

    def test_quality_xlsx_matches_filtered_csv_view_and_keeps_raw_values(self):
        page, raw = self.quality()
        before = self.snapshot()
        for code in ("el", "en", "el"):
            self.language.set_language(code, persist=False)
            self.language.apply_to(page)
            path = self.root / (code + ".xlsx")
            with patch("app.data_quality.QFileDialog.getSaveFileName", return_value=(str(path), "")), patch("app.data_quality._message"):
                page.export_xlsx()
            workbook = self.workbook(path)
            sheet = workbook.active
            rows = list(sheet.values)
            self.assertEqual(2, len(rows))
            self.assertEqual(tuple(page._issue_values(page.all_issues[0])), rows[1])
            self.assertEqual(raw, sheet["C2"].value)
            self.assertEqual("s", sheet["C2"].data_type)
            self.assertIn(raw, sheet["D2"].value)
            self.assertEqual("A2", sheet.freeze_panes)
            self.assertTrue(sheet["D2"].alignment.wrap_text)
            self.assertEqual("XLSX export" if code == "en" else "Εξαγωγή XLSX", page.xlsx_button.text())
            workbook.close()
        self.assertEqual(before, self.snapshot())

    def test_quality_cancel_and_empty_are_noops(self):
        page, _ = self.quality()
        before = set(self.root.iterdir())
        with patch("app.data_quality.QFileDialog.getSaveFileName", return_value=("", "")), \
             patch("app.data_quality.export_issue_xlsx") as writer:
            page.export_xlsx()
            writer.assert_not_called()
        page.all_issues = []
        with patch("app.data_quality.QFileDialog.getSaveFileName") as dialog, patch("app.data_quality._message"):
            page.export_xlsx()
            dialog.assert_not_called()
        self.assertEqual(before, set(self.root.iterdir()))

    def test_quality_extension_success_and_write_failure(self):
        page, _ = self.quality()
        for suffix in ("", ".XLSX"):
            path = self.root / ("issues" + suffix)
            expected = path if suffix else path.with_suffix(".xlsx")
            with patch("app.data_quality.QFileDialog.getSaveFileName", return_value=(str(path), "")), patch("app.data_quality._message") as message:
                page.export_xlsx()
            self.assertEqual(expected, message.call_args.kwargs["path"])
            original = expected.read_bytes()
            before = set(self.root.iterdir())
            with patch("app.data_quality.QFileDialog.getSaveFileName", return_value=(str(path), "")), \
                 patch("openpyxl.workbook.workbook.Workbook.save", side_effect=OSError("Raw Ω {path}")), \
                 patch("app.data_quality._message") as message:
                page.export_xlsx()
            self.assertEqual("critical", message.call_args.args[1])
            self.assertEqual("Raw Ω {path}", message.call_args.kwargs["error"])
            self.assertEqual(original, expected.read_bytes())
            self.assertEqual(before, set(self.root.iterdir()))
