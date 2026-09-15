import os
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
import csv,io,json,tempfile,unittest,subprocess,shutil
from pathlib import Path
from unittest.mock import patch
from xml.etree import ElementTree as ET
from openpyxl import load_workbook
from PySide6.QtWidgets import QApplication
from PySide6.QtGui import QFontDatabase
from app.gis.geometry import normalize,manual_coordinates
from app.gis import exports
from app.gis.importers import import_geometry
from app.database import Database
from app.gis.store import GeometryStore
from app.gis.export_dialog import CoordinateExportDialog


class CoordinateExportTests(unittest.TestCase):
    def test_cancel_and_failed_writes_preserve_destination_and_database(self):
        with tempfile.TemporaryDirectory() as folder:
            db=Database(Path(folder)/'fixture.db')
            identity=db.execute('INSERT INTO fields(name) VALUES(?)',('Export failure fixture',))
            GeometryStore(db).save(identity,self.parcels()[0].geometry)
            with db.connect() as connection:before=list(connection.iterdump())
            dialog=CoordinateExportDialog(db,identity)
            destination=Path(folder)/'existing.csv';destination.write_bytes(b'original destination')
            def partial_write(path,*args):
                Path(path).write_bytes(b'partial export')
                raise OSError('simulated storage failure')
            try:
                with patch('app.gis.export_dialog.QFileDialog.getSaveFileName',return_value=('', '')),patch('app.gis.export_dialog.exports.write') as write:
                    dialog.export();write.assert_not_called()
                for failure in ('encode','replace'):
                    with self.subTest(failure=failure),patch('app.gis.export_dialog.QFileDialog.getSaveFileName',return_value=(str(destination),'')),patch('app.gis.export_dialog.QMessageBox.warning') as warning:
                        target='app.gis.export_dialog.exports.write' if failure=='encode' else 'app.gis.export_dialog.os.replace'
                        with patch(target,side_effect=partial_write if failure=='encode' else PermissionError('destination locked')):dialog.export()
                        warning.assert_called_once()
                        self.assertEqual(b'original destination',destination.read_bytes())
                        self.assertFalse(list(Path(folder).glob('.mastixa-export-*')))
                with db.connect() as connection:self.assertEqual(before,list(connection.iterdump()))
            finally:dialog.close();dialog.deleteLater();self.qt.processEvents()

    @classmethod
    def setUpClass(cls):
        cls.qt=QApplication.instance() or QApplication([])
        font=Path('C:/Windows/Fonts/segoeui.ttf')
        if font.exists():QFontDatabase.addApplicationFont(str(font))

    def parcels(self):
        geometry=normalize({'type':'MultiPolygon','coordinates':[
            [[[26,38],[26.01,38],[26.01,38.01],[26,38.01],[26,38]],[[26.002,38.002],[26.002,38.004],[26.004,38.004],[26.004,38.002],[26.002,38.002]]],
            [[[26.02,38],[26.021,38],[26.02,38.001],[26.02,38]]]]},'EPSG:4326')
        return [exports.ExportParcel('fixture-id','=Χωράφι / Α','001234','manual',geometry)]

    def test_all_formats_preserve_order_holes_unicode_and_numeric_values(self):
        parcels=self.parcels();rows=exports.table(parcels,'wgs84');self.assertEqual(12,len(rows));self.assertEqual(['Latitude','Longitude'],rows[0][-2:])
        self.assertEqual([1,1,1],rows[1][7:10]);self.assertEqual([1,2,1],rows[5][7:10]);self.assertEqual([2,1,1],rows[9][7:10]);self.assertEqual(38,rows[1][10]);self.assertEqual(26,rows[1][11])
        csvrows=list(csv.reader(io.StringIO(exports.csv_bytes(rows).decode('utf-8-sig')),delimiter=';'))
        self.assertEqual("'=Χωράφι / Α",csvrows[1][0]);self.assertEqual('001234',csvrows[1][1]);self.assertEqual('38.00000000',csvrows[1][10])
        book=load_workbook(io.BytesIO(exports.xlsx_bytes(rows)),data_only=False);sheet=book['Όλες οι κορυφές']
        self.assertEqual('=Χωράφι / Α',sheet['A2'].value);self.assertEqual('s',sheet['A2'].data_type);self.assertEqual('n',sheet['K2'].data_type);self.assertEqual(38,sheet['K2'].value);self.assertEqual('001234',sheet['B2'].value);self.assertEqual(12,sheet.max_row)
        geojson=json.loads(exports.geojson_bytes(parcels));feature=geojson['features'][0]
        self.assertEqual(parcels[0].geometry.wgs84,feature['geometry']);self.assertIn('perimeter_m',feature['properties']);self.assertNotIn('perimeter_m2',feature['properties'])
        restored=import_geometry(exports.geojson_bytes(parcels),'fixture.geojson','');self.assertEqual(11,restored.vertex_count)
        kml=exports.kml_bytes(parcels);root=ET.fromstring(kml);ns={'k':'http://www.opengis.net/kml/2.2'}
        self.assertEqual(2,len(root.findall('.//k:Polygon',ns)));self.assertEqual(1,len(root.findall('.//k:innerBoundaryIs',ns)));self.assertEqual(parcels[0].name,root.find('.//k:name',ns).text)
        self.assertEqual(11,import_geometry(kml,'fixture.kml','').vertex_count)
        with tempfile.TemporaryDirectory() as folder:
            pdf=Path(folder)/'vertices.pdf';exports.write(pdf,parcels,'wgs84','pdf');self.assertTrue(pdf.read_bytes().startswith(b'%PDF'))
            if shutil.which('pdftotext'):
                # Xpdf (Git for Windows) defaults to non-UTF-8 output; align producer and decoder.
                text=subprocess.run(['pdftotext','-enc','UTF-8','-layout',str(pdf),'-'],check=True,capture_output=True,encoding='utf-8').stdout
                self.assertIn('Χωράφι',text);self.assertIn('EPSG:4326',text);self.assertIn('38.00000000',text);self.assertIn('26.02000000',text)

    def test_multifield_source_precision_negative_coordinates_and_names(self):
        parcels=self.parcels();source=manual_coordinates('500000.12345678 4200000.12345678\n500010.12345678 4200000.12345678\n500010.12345678 4200010.12345678\n500000.12345678 4200010.12345678','EPSG:2100')
        parcels.append(exports.ExportParcel('projected','ΕΓΣΑ87','002','source',source));rows=exports.table(parcels,'source')
        self.assertEqual(16,len(rows));self.assertEqual('EPSG:2100',rows[-1][3]);self.assertEqual('500000.12345678',str(rows[-1][10]));self.assertEqual('4200010.12345678',str(rows[-1][11]))
        west=exports.ExportParcel('west','West','003','manual',manual_coordinates('-3 38\n-2.999 38\n-2.999 38.001\n-3 38.001','EPSG:4326'))
        westcsv=list(csv.reader(io.StringIO(exports.csv_bytes(exports.table([west],'wgs84')).decode('utf-8-sig')),delimiter=';'))
        self.assertEqual('-3.00000000',westcsv[1][11]);self.assertEqual(2,len(json.loads(exports.geojson_bytes(parcels))['features']))
        self.assertNotIn('/',exports.filename(parcels[:1],'source','xlsx'));self.assertIn('2_fields',exports.filename(parcels,'wgs84','csv'))

    def test_export_selection_preview_does_not_change_data(self):
        with tempfile.TemporaryDirectory() as folder:
            db=Database(Path(folder)/'fixture.db');ids=[];store=GeometryStore(db)
            for name in ['Πρώτο','Δεύτερο','Χωρίς όρια']:
                ids.append(db.execute('INSERT INTO fields(name) VALUES(?)',(name,)))
            for identity in ids[:2]:store.save(identity,manual_coordinates('26 38\n26.001 38\n26.001 38.001\n26 38.001','EPSG:4326'))
            before=db.query_one('SELECT count(*) AS n FROM gis_changes')['n'];dialog=CoordinateExportDialog(db,ids[0])
            try:
                self.assertEqual(1,len(dialog.parcels));dialog.select_all(True);self.assertEqual(2,len(dialog.parcels));self.assertEqual(5,dialog.preview.rowCount());dialog.mode.setCurrentIndex(1);self.assertEqual('source',dialog.effective_mode());dialog.format.setCurrentIndex(3);self.assertEqual('wgs84',dialog.effective_mode());self.assertFalse(dialog.mode.isEnabled());dialog.select_all(False);self.assertFalse(dialog.parcels);self.assertEqual(before,db.query_one('SELECT count(*) AS n FROM gis_changes')['n'])
            finally:dialog.close();dialog.deleteLater();self.qt.processEvents()


if __name__=='__main__':unittest.main()
